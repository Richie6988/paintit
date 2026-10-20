"""Pont Excel avec l'ERP d'un fournisseur (MangoERP / 芒果店长, ou tout ERP qui importe des commandes
par fichier), SANS Shopify.

SORTANT : le fournisseur cree dans son ERP un « magasin personnalise » (自定义店铺) puis importe nos
  commandes via « Excel导入手工订单 ». On remplit SON modele Excel (televerse sur sa fiche) : chaque
  colonne du modele est associee a un champ PaintIt (detection automatique FR/EN/中文, ajustable).
ENTRANT : le fournisseur exporte ses commandes expediees (Excel/CSV) ; on lit n° de commande +
  transporteur + n° de suivi et on met a jour nos commandes (meme traitement que le webhook).
"""
import csv
import io
import re
import unicodedata

from django.conf import settings

# (cle, libelle, mots-cles de detection dans l'en-tete du modele)
FIELDS = [
    ("uid", "N° de commande PaintIt", ["订单号", "订单编号", "平台订单", "order no", "order number", "order id", "commande", "参考号", "reference"]),
    ("date", "Date de commande", ["日期", "下单时间", "付款时间", "date", "time"]),
    ("sku_qty", "SKU*quantité (MangoERP : sku*1;sku*1)", ["sku编码和数量", "sku and qty", "sku*"]),
    ("sku", "SKU (modèle ci-dessous)", ["sku", "商品编码", "产品编码", "货号", "seller sku"]),
    ("amount", "Montant (prix d'achat fournisseur, pas le prix de vente)", ["订单金额", "订单总额", "order amount", "order total", "montant"]),
    ("product", "Nom du produit", ["商品名称", "产品名称", "品名", "标题", "title", "product name", "product", "item name"]),
    ("variant", "Variante (format · couleurs)", ["属性", "规格", "变体", "variant", "variation", "option"]),
    ("qty", "Quantité (1)", ["数量", "qty", "quantity", "quantité"]),
    ("unit_cost", "Prix unitaire (coût fournisseur)", ["单价", "价格", "售价", "unit price", "price", "prix"]),
    ("currency", "Devise (EUR)", ["币种", "货币", "currency", "devise"]),
    ("name", "Destinataire", ["收件人", "收货人", "买家姓名", "姓名", "recipient", "receiver", "consignee", "name", "nom"]),
    ("phone", "Téléphone", ["电话", "联系电话", "phone", "tel", "téléphone"]),
    ("mobile", "Mobile (même numéro)", ["手机", "mobile", "portable"]),
    ("email", "E-mail", ["邮箱", "email", "e-mail", "mail"]),
    ("address_full", "Adresse complète (ligne 1 + 2)", ["详细地址", "full address", "adresse complète"]),
    ("address1", "Adresse ligne 1", ["地址1", "街道1", "address1", "address 1", "address line 1", "street", "地址"]),
    ("address2", "Adresse ligne 2", ["地址2", "街道2", "address2", "address 2", "address line 2", "complément"]),
    ("postal_code", "Code postal", ["邮编", "邮政编码", "zip", "postal", "postcode", "code postal"]),
    ("city", "Ville", ["城市", "city", "ville"]),
    ("state", "Région / État", ["省", "州", "province", "state", "region", "région"]),
    ("country", "Pays (nom)", ["国家名", "country name", "pays", "country"]),
    ("country_code", "Pays (code ISO2)", ["国家二字码", "国家代码", "国家简码", "二字码", "country code", "countrycode", "iso"]),
    ("ioss", "N° IOSS (réglage)", ["ioss"]),
    ("sender_tax", "Votre n° de TVA (Légal)", ["寄件人税号", "sender tax", "vat"]),
    ("logistics", "Transporteur choisi (réglage)", ["买家所选物流", "物流方式", "shipping method", "logistics"]),
    ("note", "Note (fichiers de production)", ["备注", "买家留言", "留言", "remark", "note", "memo", "comment"]),
    ("files", "Lien ZIP des fichiers de production", ["图片", "链接", "附件", "image", "file", "url", "link", "fichier"]),
]
FIELD_LABELS = dict((k, l) for k, l, _ in FIELDS)

# Colonnes du fichier de SUIVI renvoye par le fournisseur
TRACK_KEYS = {
    "uid": ["订单号", "订单编号", "平台订单", "参考号", "order no", "order number", "order id", "commande", "reference"],
    "tracking_number": ["跟踪号", "追踪号", "运单号", "物流单号", "tracking", "waybill", "suivi"],
    "carrier": ["物流方式", "物流渠道", "承运商", "物流公司", "carrier", "shipping method", "logistics", "transporteur"],
}

DEFAULT_SKU = "PBN-{format}-{colors}C"


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode() or str(s or "")
    return re.sub(r"[\s_*:：()（）\-]+", " ", s.lower()).strip()


def _match(header, keywords):
    h, raw = _norm(header), str(header or "").lower()
    for k in keywords:
        if k in raw or (_norm(k) and _norm(k) in h):
            return len(k)
    return 0


def auto_map(headers):
    """{index_colonne: cle_champ} : chaque colonne recoit le champ dont le mot-cle correspond le mieux."""
    out, used = {}, set()
    for i, h in enumerate(headers):
        best, score = "", 0
        for key, _l, kws in FIELDS:
            sc = _match(h, kws)
            if sc > score and key not in used:
                best, score = key, sc
        if best:
            out[i] = best; used.add(best)
    return out


def read_template(path):
    """-> (headers, header_row) : 1re ligne (parmi les 10 premieres) avec au moins 3 cellules remplies."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=10, values_only=True), start=1):
        vals = ["" if v is None else str(v).strip() for v in row]
        if sum(1 for v in vals if v) >= 3:
            while vals and not vals[-1]:
                vals.pop()
            return vals, r_idx
    return [], 1


def order_values(sup, order, sku_pattern=None):
    """Valeurs d'une commande (modele Order) pour chaque champ PaintIt."""
    from . import address
    from .security import file_url
    from .models import format_key
    fmt = format_key(order.width_cm, order.height_cm)
    pattern = sku_pattern or DEFAULT_SKU
    try:
        sku = pattern.format(format=fmt, colors=order.colors, brushes="B" if order.brushes else "")
    except (KeyError, IndexError, ValueError):
        sku = DEFAULT_SKU.format(format=fmt, colors=order.colors)
    cost = sup.goods_cost(fmt, order.colors, order.brushes)
    # Meme contenu que si la commande arrivait de la boutique Shopify : SKU, titre produit et variante,
    # telephone international, proprietes de ligne (fichiers) -> le fournisseur traite les deux a l'identique.
    from . import shopify as _shop
    from .models import CompanyInfo
    variant = _shop.variant_title(order) + (" + Set de pinceaux" if order.brushes else "")
    cfg = (sup.offer or {}).get("xlsx") or {}
    phone = _shop._phone(order.phone_code, order.phone) or ""
    sku_qty = "%s*1" % sku + (";%s*1" % (cfg.get("brush_sku") or _shop.BRUSH_SKU) if order.brushes else "")
    props = _shop.line_properties(order)
    return {
        "sku_qty": sku_qty, "amount": _shop.supplier_amounts(order, sup)[2], "mobile": phone,
        "address_full": ", ".join(x for x in (order.address1, order.address2) if x),
        "ioss": cfg.get("ioss", ""), "logistics": cfg.get("logistics", ""),
        "sender_tax": getattr(CompanyInfo.get(), "vat_number", "") or "",
        "uid": order.uid, "date": order.created_at.strftime("%Y-%m-%d %H:%M") if order.created_at else "",
        "sku": sku, "product": _shop.PRODUCT_TITLE, "variant": variant,
        "qty": 1, "unit_cost": round(cost, 2) if cost is not None else "", "currency": "EUR",
        "name": order.customer_name, "phone": phone,
        "email": order.customer_email, "address1": order.address1, "address2": order.address2,
        "postal_code": order.postal_code, "city": order.city, "state": order.city or "",   # province obligatoire chez certains ERP
        "country": order.country, "country_code": (address.COUNTRIES.get(order.country) or ("",))[0],
        # proprietes de ligne Shopify, une par ligne (« Nom: valeur ») : Mango n'accepte pas de piece jointe a l'import
        "note": "\n".join("%s: %s" % (n, v) for n, v in props)[:2000],
        "files": __import__("studio.security", fromlist=["x"]).production_url(order.uid),
    }


def build_xlsx(sup, orders):
    """Remplit le modele du fournisseur (ou un modele standard) avec les commandes -> bytes .xlsx."""
    import openpyxl
    cfg = (sup.offer or {}).get("xlsx") or {}
    tpl = cfg.get("template")
    mapping = {int(k): v for k, v in (cfg.get("map") or {}).items() if v}
    path = None
    if tpl:
        import os
        path = os.path.join(settings.MEDIA_ROOT, tpl)
        if not os.path.exists(path):
            path = None
    if path:
        wb = openpyxl.load_workbook(path)
        ws = wb.worksheets[0]
        start = int(cfg.get("header_row") or 1) + 1
        # on repart d'un modele vide : lignes d'exemple eventuelles effacees
        if ws.max_row >= start:
            ws.delete_rows(start, ws.max_row - start + 1)
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "orders"
        keys = [k for k, _l, _ in FIELDS]
        ws.append(["%s / %s" % (FIELD_LABELS[k], k) for k in keys])
        mapping = dict(enumerate(keys))
        start = 2
    for r, o in enumerate(orders, start=start):
        vals = order_values(sup, o, cfg.get("sku"))
        for col, key in mapping.items():
            ws.cell(row=r, column=col + 1, value=vals.get(key, ""))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_rows(upload_name, data):
    """Lit un fichier de suivi (.xlsx ou .csv) -> (headers, rows)."""
    if upload_name.lower().endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows = [["" if v is None else str(v).strip() for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
    else:
        text = None
        for enc in ("utf-8-sig", "gb18030", "latin-1"):
            try:
                text = data.decode(enc); break
            except UnicodeDecodeError:
                continue
        dialect = csv.Sniffer().sniff(text[:2000], delimiters=",;\t") if text else csv.excel
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text or ""), dialect)]
    rows = [r for r in rows if any(r)]
    for i, r in enumerate(rows[:10]):
        if sum(1 for c in r if c) >= 2 and any(_match(c, TRACK_KEYS["uid"]) for c in r):
            return r, rows[i + 1:]
    return (rows[0], rows[1:]) if rows else ([], [])


def import_tracking(sup, upload_name, data):
    """Met a jour les commandes a partir du fichier de suivi -> (n_maj, erreurs)."""
    from .integrations import handle_event
    headers, rows = read_rows(upload_name, data)
    col = {}
    for key, kws in TRACK_KEYS.items():
        best, score = None, 0
        for i, h in enumerate(headers):
            sc = _match(h, kws)
            if sc > score:
                best, score = i, sc
        if best is not None:
            col[key] = best
    if "uid" not in col or "tracking_number" not in col:
        return 0, ["Colonnes introuvables : il faut au moins « n° de commande » (订单号) et « n° de suivi » (跟踪号)."]
    n, errors = 0, []
    for r in rows:
        get = lambda k: (r[col[k]] if k in col and col[k] < len(r) else "").strip()
        uid, trk = get("uid"), get("tracking_number")
        if not uid or not trk:
            continue
        code, msg = handle_event(sup, {"uid": uid, "status": "shipped", "tracking_number": trk,
                                       "carrier": get("carrier"), "message": "import fichier de suivi"})
        if code == 200:
            n += 1
        else:
            errors.append("%s : %s" % (uid, msg))
    return n, errors
